from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "program" / "final"

SOFTWARE_NAME = "教学资源智能采集与知识库构建系统软件"
VERSION = "V1.0"
MATERIAL_TYPE = "程序鉴别材料"
TITLE = f"{SOFTWARE_NAME}_{VERSION}_{MATERIAL_TYPE}"
HEADER_TEXT = f"{SOFTWARE_NAME} {VERSION} {MATERIAL_TYPE}"
FOOTER_TEXT = "著作权人名称（提交前替换为与申请表一致的真实全称）"

DOCX_PATH = OUTPUT_DIR / f"{TITLE}.docx"
PDF_PATH = OUTPUT_DIR / f"{TITLE}.pdf"
MANIFEST_PATH = OUTPUT_DIR / f"{TITLE}_源码清单.txt"

PAGE_LINE_CAPACITY = 60
MIN_CODE_LINES_PER_PAGE = 50
MAX_CODE_CHARS = 118


# The list intentionally includes only self-developed product source files.
# Excluded: virtual environments, dependencies, documents, logs, data, generated
# outputs, screenshots, model results, tests, debugging scripts, and this
# copyright material generator.
SOURCE_FILES = [
    # Startup, Web console, and task orchestration.
    "start_web.py",
    "web/app.py",
    "web/services.py",
    "web/templates/index.html",
    "web/static/script.js",
    "web/static/style.css",
    "api/app.py",
    "api/auth.py",
    "api/v1/admin.py",
    "api/v1/search.py",
    "server.py",
    "client.py",
    "registry.py",
    "llm_client.py",
    "llm_extractor.py",
    "exporter.py",
    "exporter_v2.py",
    # Teaching resource collection, page fetching, and text cleaning.
    "utils/env_loader.py",
    "utils/crawler_framework.py",
    "utils/webpage_crawler.py",
    "utils/playwright_fetcher.py",
    "utils/html_cleaner.py",
    "utils/image_analyzer.py",
    "utils/webpage_storage.py",
    "utils/web_deduplication.py",
    # Knowledge type recognition and structured extraction.
    "extractors/se_kb_types.py",
    "extractors/type_registry.py",
    "extractors/type_router.py",
    "extractors/specialized_extractors.py",
    "extractors/se_kb_extractors.py",
    "extractors/pipeline.py",
    "adapters/__init__.py",
    "adapters/concepts_adapter.py",
    "adapters/dfd_adapter.py",
    "adapters/rules_adapter.py",
    "adapters/enhanced_adapters.py",
    "adapters/llm_preselector.py",
    "validators/jsonschema_validator.py",
    "validators/semantic_stage3.py",
    # Command tools and processing pipelines.
    "scripts/manual_crawl.py",
    "scripts/run_web_crawl.py",
    "scripts/universal_knowledge_processor.py",
    "scripts/format_processor.py",
    "scripts/evaluate_quality.py",
    "scripts/build_vector_index.py",
    "scripts/query_knowledge.py",
    "scripts/refine_knowledge.py",
    # Result writing, deduplication/refinement, semantic retrieval, and config.
    "writers/artifacts_writer.py",
    "refiner/deduplicator.py",
    "refiner/embedder.py",
    "refiner/merger.py",
    "refiner/knowledge_refiner.py",
    "vectorizer/config.py",
    "vectorizer/chunker.py",
    "vectorizer/store.py",
    "vectorizer/indexer.py",
    "vectorizer/query_planner.py",
    "vectorizer/retriever.py",
    "config/path_config.py",
]


@dataclass
class Block:
    lines: list[str]
    code_display_lines: int
    file_path: str | None = None
    is_file_marker: bool = False


@dataclass
class Page:
    blocks: list[Block]

    @property
    def display_line_count(self) -> int:
        return sum(len(block.lines) for block in self.blocks)

    @property
    def code_display_line_count(self) -> int:
        return sum(block.code_display_lines for block in self.blocks)

    @property
    def first_file_path(self) -> str | None:
        for block in self.blocks:
            if block.file_path:
                return block.file_path
        return None

    def rendered_lines(self) -> list[str]:
        rendered: list[str] = []
        for block in self.blocks:
            rendered.extend(block.lines)
        return rendered


def read_source(path: Path) -> list[str]:
    return path.read_text(encoding="utf-8", errors="replace").splitlines()


def wrap_code_line(code: str) -> list[str]:
    expanded = code.expandtabs(4).rstrip()
    if expanded == "":
        return [""]

    wrapped: list[str] = []
    remaining = expanded
    while len(remaining) > MAX_CODE_CHARS:
        cut = MAX_CODE_CHARS
        window = remaining[: MAX_CODE_CHARS + 1]
        candidates = [
            window.rfind(" "),
            window.rfind(","),
            window.rfind(")"),
            window.rfind("]"),
            window.rfind("}"),
        ]
        best = max(candidates)
        if best >= 72:
            cut = best + 1
        wrapped.append(remaining[:cut].rstrip())
        remaining = remaining[cut:].lstrip()
    wrapped.append(remaining)
    return wrapped


def make_blocks(source_files: Iterable[str]) -> tuple[list[Block], int, list[tuple[str, int]]]:
    blocks: list[Block] = []
    global_line_no = 1
    total_source_lines = 0
    file_stats: list[tuple[str, int]] = []

    for relative_path in source_files:
        path = ROOT / relative_path
        if not path.exists():
            raise FileNotFoundError(relative_path)

        lines = read_source(path)
        total_source_lines += len(lines)
        file_stats.append((relative_path, len(lines)))
        blocks.append(
            Block([f"# 文件：{relative_path}"], code_display_lines=0, file_path=relative_path, is_file_marker=True)
        )

        for line in lines:
            wrapped = wrap_code_line(line)
            rendered = [f"{global_line_no:06d} {wrapped[0]}"]
            rendered.extend(f"       {part}" for part in wrapped[1:])
            blocks.append(Block(rendered, code_display_lines=len(rendered), file_path=relative_path))
            global_line_no += 1

    return blocks, total_source_lines, file_stats


def page_marker(path: str, continued: bool = True) -> Block:
    suffix = "（续）" if continued else ""
    return Block([f"# 文件：{path}{suffix}"], code_display_lines=0, file_path=path, is_file_marker=True)


def paginate_blocks(blocks: list[Block], metadata_lines: list[str]) -> list[Page]:
    pages: list[Page] = []
    current: list[Block] = [Block(metadata_lines, code_display_lines=0, file_path=None)] if metadata_lines else []
    current_lines = sum(len(block.lines) for block in current)
    current_file: str | None = None

    def flush(next_file: str | None = None) -> None:
        nonlocal current, current_lines, current_file
        if current:
            pages.append(Page(current))
        current = []
        current_lines = 0
        current_file = None
        if next_file:
            marker = page_marker(next_file)
            current = [marker]
            current_lines = len(marker.lines)
            current_file = next_file

    for block in blocks:
        block_len = len(block.lines)
        if block_len > PAGE_LINE_CAPACITY:
            raise ValueError(f"Single block is too large to fit a page: {block.file_path}")

        starts_mid_file = (
            not current
            and block.file_path is not None
            and not block.is_file_marker
        )
        if starts_mid_file:
            marker = page_marker(block.file_path)
            current.append(marker)
            current_lines += len(marker.lines)
            current_file = block.file_path

        if current_lines + block_len > PAGE_LINE_CAPACITY:
            flush(block.file_path if block.file_path and not block.is_file_marker else None)
            if current_lines + block_len > PAGE_LINE_CAPACITY:
                flush(None)

        if block.is_file_marker and current and current[-1].is_file_marker:
            # Avoid adjacent file marker lines when a new page was opened with a
            # continuation marker and the next block is already a file marker.
            current[-1] = block
        else:
            current.append(block)
            current_lines += block_len
        if block.file_path:
            current_file = block.file_path

    if current:
        pages.append(Page(current))

    rebalance_last_page(pages)
    return pages


def rebalance_last_page(pages: list[Page]) -> None:
    if len(pages) < 2:
        return
    last = pages[-1]
    prev = pages[-2]
    while (
        last.code_display_line_count < MIN_CODE_LINES_PER_PAGE
        and prev.blocks
        and prev.code_display_line_count > MIN_CODE_LINES_PER_PAGE
        and last.display_line_count < PAGE_LINE_CAPACITY
    ):
        moved = prev.blocks.pop()
        if moved.is_file_marker:
            continue
        insert_at = 1 if last.blocks and last.blocks[0].is_file_marker else 0
        last.blocks.insert(insert_at, moved)
        if last.first_file_path and (not last.blocks[0].is_file_marker):
            marker = page_marker(last.first_file_path)
            last.blocks.insert(0, marker)
        if prev.blocks and prev.blocks[-1].is_file_marker:
            prev.blocks.pop()


def select_submission_pages(pages: list[Page]) -> tuple[list[Page], str]:
    if len(pages) <= 60:
        return pages, "全部源程序"
    return pages[:30] + pages[-30:], "源程序排版后超过 60 页，提交前连续 30 页和后连续 30 页"


def replace_metadata_placeholders(metadata_lines: list[str], *, total_pages: int, mode: str) -> list[str]:
    result = []
    for line in metadata_lines:
        line = line.replace("{TOTAL_PAGES}", str(total_pages)).replace("{MODE}", mode)
        result.append(line)
    return result


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(f"{FOOTER_TEXT}    第 ")
    run = paragraph.add_run()

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_sep)
    run._r.append(fld_char_end)
    paragraph.add_run(" 页")


def set_run_font(run, name: str, east_asia: str | None = None, size_pt: float | None = None) -> None:
    run.font.name = name
    if size_pt:
        run.font.size = Pt(size_pt)
    r_fonts = run._element.rPr.rFonts
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east_asia or name)


def render_docx(pages: list[Page], docx_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(1.05)
    section.left_margin = Cm(1.05)
    section.right_margin = Cm(1.05)
    section.header_distance = Cm(0.45)
    section.footer_distance = Cm(0.45)

    header_paragraph = section.header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header_paragraph.add_run(HEADER_TEXT)
    set_run_font(header_run, "Microsoft YaHei", "微软雅黑", 8.5)

    footer_paragraph = section.footer.paragraphs[0]
    add_page_number(footer_paragraph)
    for run in footer_paragraph.runs:
        set_run_font(run, "Microsoft YaHei", "微软雅黑", 8.0)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Consolas"
    normal.font.size = Pt(7.0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for page_index, page in enumerate(pages):
        for line in page.rendered_lines():
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            paragraph.paragraph_format.line_spacing = Pt(10.2)
            run = paragraph.add_run(line)
            set_run_font(run, "Consolas", "宋体", 7.0)
        if page_index != len(pages) - 1:
            document.paragraphs[-1].runs[-1].add_break(WD_BREAK.PAGE)

    document.save(docx_path)


def register_pdf_font() -> str:
    font_candidates = [
        ("ProgramCodeFont", Path("C:/Windows/Fonts/simsun.ttc")),
        ("ProgramCodeFont", Path("C:/Windows/Fonts/simhei.ttf")),
        ("ProgramCodeFont", Path("C:/Windows/Fonts/msyh.ttc")),
    ]
    for font_name, path in font_candidates:
        if path.exists():
            try:
                pdfmetrics.registerFont(TTFont(font_name, str(path)))
                return font_name
            except Exception:
                continue
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def render_pdf(pages: list[Page], pdf_path: Path) -> None:
    font_name = register_pdf_font()
    width, height = A4
    c = canvas.Canvas(str(pdf_path), pagesize=A4)

    left_margin = 30
    top_y = height - 50
    line_height = 11.2

    for page_no, page in enumerate(pages, start=1):
        c.setFont(font_name, 8.5)
        c.drawCentredString(width / 2, height - 24, HEADER_TEXT)
        c.setFont(font_name, 8.0)
        c.drawCentredString(width / 2, 24, f"{FOOTER_TEXT}    第 {page_no} 页")

        c.setFont(font_name, 7.0)
        y = top_y
        for line in page.rendered_lines():
            safe_line = line.replace("\t", "    ")
            c.drawString(left_margin, y, safe_line)
            y -= line_height
        c.showPage()

    c.save()


def write_manifest(
    manifest_path: Path,
    file_stats: list[tuple[str, int]],
    *,
    total_source_lines: int,
    full_pages: int,
    submission_pages: int,
    mode: str,
    min_page_code_lines: int,
) -> None:
    lines = [
        f"软件名称：{SOFTWARE_NAME}",
        f"版本号：{VERSION}",
        f"材料类型：{MATERIAL_TYPE}",
        f"纳入源码文件数：{len(file_stats)}",
        f"自主开发源程序总行数：{total_source_lines}",
        f"完整排版页数：{full_pages}",
        f"提交页数：{submission_pages}",
        f"提交范围：{mode}",
        f"提交页每页最少代码显示行数：{min_page_code_lines}",
        "",
        "纳入源码清单：",
    ]
    lines.extend(f"{path}\t{line_count}" for path, line_count in file_stats)
    manifest_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def validate_pages(pages: list[Page]) -> None:
    low_pages = [
        (index + 1, page.code_display_line_count)
        for index, page in enumerate(pages)
        if page.code_display_line_count < MIN_CODE_LINES_PER_PAGE
    ]
    if low_pages:
        raise ValueError(f"Pages below {MIN_CODE_LINES_PER_PAGE} code display lines: {low_pages[:10]}")


def validate_output_text() -> None:
    forbidden_in_formal_metadata = [
        "SE-KB-Engine" + "知识库构建系统软件 " + "V0.1.0",
        "SE-KB-Engine " + "V0.1.0",
        "V0.1.0 " + "程序鉴别材料",
    ]
    for text in [TITLE, HEADER_TEXT, FOOTER_TEXT]:
        for forbidden in forbidden_in_formal_metadata:
            if forbidden in text:
                raise ValueError(f"Forbidden legacy formal metadata found: {forbidden}")


def main() -> None:
    validate_output_text()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    blocks, total_source_lines, file_stats = make_blocks(SOURCE_FILES)
    metadata_template = [
        f"# 软件名称：{SOFTWARE_NAME}",
        f"# 版本号：{VERSION}",
        f"# 材料类型：{MATERIAL_TYPE}",
        f"# 著作权人：{FOOTER_TEXT}",
        f"# 生成范围：纳入自主开发源码 {len(file_stats)} 个文件，共 {total_source_lines} 行；完整排版 {{TOTAL_PAGES}} 页；{{MODE}}。",
        "# 以下为项目真实源程序内容，按功能链路排序并保留源码原始路径标识。",
    ]
    preliminary_pages = paginate_blocks(blocks, metadata_template)
    submission_pages, mode = select_submission_pages(preliminary_pages)
    metadata = replace_metadata_placeholders(metadata_template, total_pages=len(preliminary_pages), mode=mode)
    final_pages = paginate_blocks(blocks, metadata)
    submission_pages, mode = select_submission_pages(final_pages)
    validate_pages(submission_pages)

    min_code_lines = min(page.code_display_line_count for page in submission_pages)
    render_docx(submission_pages, DOCX_PATH)
    render_pdf(submission_pages, PDF_PATH)
    write_manifest(
        MANIFEST_PATH,
        file_stats,
        total_source_lines=total_source_lines,
        full_pages=len(final_pages),
        submission_pages=len(submission_pages),
        mode=mode,
        min_page_code_lines=min_code_lines,
    )

    print(f"DOCX: {DOCX_PATH}")
    print(f"PDF: {PDF_PATH}")
    print(f"Manifest: {MANIFEST_PATH}")
    print(f"Source files: {len(file_stats)}")
    print(f"Source lines: {total_source_lines}")
    print(f"Full pages: {len(final_pages)}")
    print(f"Submission pages: {len(submission_pages)}")
    print(f"Mode: {mode}")
    print(f"Min code display lines/page: {min_code_lines}")


if __name__ == "__main__":
    main()

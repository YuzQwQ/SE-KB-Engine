# -*- coding: utf-8 -*-
"""
Generate the formal software copyright design document.

Outputs:
- docs/软著鉴别材料_教学资源智能采集与知识库构建系统软件_V1.0_软件设计说明书.md
- output/doc/教学资源智能采集与知识库构建系统软件_V1.0_软件设计说明书.md
- output/doc/教学资源智能采集与知识库构建系统软件_V1.0_软件设计说明书.docx
- output/doc/教学资源智能采集与知识库构建系统软件_V1.0_软件设计说明书.pdf
- docs/软著提交内部检查清单_教学资源智能采集与知识库构建系统软件_V1.0.md
"""

from __future__ import annotations

import re
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Literal

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepTogether,
    Image as PdfImage,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parent.parent
DOCS_DIR = ROOT / "docs"
OUTPUT_DIR = ROOT / "output" / "doc"

SOFTWARE_NAME = "教学资源智能采集与知识库构建系统软件"
VERSION = "V1.0"
INTERNAL_CODE_NAME = "SE-KB-Engine"
FULL_NAME = f"{SOFTWARE_NAME} {VERSION}"
DOC_TITLE = f"{FULL_NAME} 软件设计说明书"
COPYRIGHT_OWNER = os.getenv("COPYRIGHT_OWNER", "任政昆 陈锟烨 尤志宁").strip()
COPYRIGHT_OWNER_FORMAL = (
    COPYRIGHT_OWNER
    if COPYRIGHT_OWNER != "以申请表登记信息为准"
    else "以《计算机软件著作权登记申请表》登记信息为准"
)
DOC_SOURCE = DOCS_DIR / "软著鉴别材料_教学资源智能采集与知识库构建系统软件_V1.0_软件设计说明书.md"
OUTPUT_BASENAME = "教学资源智能采集与知识库构建系统软件_V1.0_软件设计说明书"
INTERNAL_CHECKLIST = DOCS_DIR / "软著提交内部检查清单_教学资源智能采集与知识库构建系统软件_V1.0.md"


ElementKind = Literal["heading", "paragraph", "bullets", "table", "code", "image", "pagebreak"]


@dataclass(frozen=True)
class Element:
    kind: ElementKind
    text: str = ""
    level: int = 1
    items: tuple[str, ...] = ()
    headers: tuple[str, ...] = ()
    rows: tuple[tuple[str, ...], ...] = ()
    widths: tuple[float, ...] = ()
    caption: str = ""


def H(level: int, text: str) -> Element:
    return Element("heading", text=text, level=level)


def P(text: str) -> Element:
    return Element("paragraph", text=text)


def B(*items: str) -> Element:
    return Element("bullets", items=items)


def T(headers: Iterable[str], rows: Iterable[Iterable[str]], widths: Iterable[float] = ()) -> Element:
    return Element(
        "table",
        headers=tuple(headers),
        rows=tuple(tuple(row) for row in rows),
        widths=tuple(widths),
    )


def C(text: str) -> Element:
    return Element("code", text=text)


def IMG(path: str, caption: str) -> Element:
    return Element("image", text=path, caption=caption)


def PB() -> Element:
    return Element("pagebreak")


def resolve_asset(path_text: str) -> Path:
    path = Path(path_text)
    return path if path.is_absolute() else ROOT / path


ELEMENTS: list[Element] = [
    H(1, DOC_TITLE),
    T(
        ("项目", "内容"),
        (
            ("软件名称", SOFTWARE_NAME),
            ("版本号", VERSION),
            ("文档类型", "软件设计说明书"),
            ("适用材料", "计算机软件著作权登记鉴别材料"),
            ("著作权人", COPYRIGHT_OWNER_FORMAL),
        ),
        (4.0, 11.8),
    ),
    P(
        "本文档用于说明本软件的设计目标、运行环境、总体结构、功能模块、数据流程、"
        "数据组织方式、关键处理流程和基本使用方式。本文档中的软件名称、版本号、"
        "页眉页脚和输出文件名均采用正式申报口径。"
    ),
    H(2, "目录"),
    B(
        "第一章 软件概述",
        "第二章 运行环境与配置",
        "第三章 系统总体架构",
        "第四章 功能模块设计",
        "第五章 数据流程设计",
        "第六章 数据组织与存储设计",
        "第七章 关键处理流程设计",
        "第八章 异常处理与运行保障",
        "第九章 使用说明",
        "第十章 维护、安全与版本管理",
        "附录 A 术语说明",
        "附录 B 登记信息一致性说明",
        "附录 C 主要配置项说明",
        "附录 D 主要成果文件示例说明",
    ),
    H(1, "第一章 软件概述"),
    H(2, "1.1 软件名称与版本"),
    P(
        f"{FULL_NAME}是一款面向高校课程、教学平台、公开教学网页和课程知识沉淀场景的"
        "教学资源处理软件。软件通过浏览器提供 Web 控制台，用户可在本地或服务器环境中"
        "完成公开教学资源采集、网页正文获取、课程知识整理、结构化知识抽取、去重精炼、"
        "语义检索和教学知识库服务等操作。"
    ),
    H(2, "1.2 软件名称与工程代号说明"),
    P(
        f"本软件正式申报名称为“{FULL_NAME}”；{INTERNAL_CODE_NAME} 为软件开发过程中的内部工程代号，"
        "可能出现在代码目录、配置文件、运行日志或成果目录中；二者指向同一套软件系统。"
        "本文档除解释内部工程代号时外，统一使用正式申报名称或“本软件”进行表述。"
    ),
    P(
        "本软件在浏览器页面中可能显示界面名称“知识库构建系统”，该名称仅作为界面标题或操作入口说明，"
        "不作为软件著作权登记的正式申报名称。"
    ),
    H(2, "1.3 设计目标"),
    P(
        "本软件的设计目标是将分散在公开网页和教学平台中的课程资料、软件工程教学资料、"
        "公开教学网页和课程知识内容整理为可核对、可复用、可检索的结构化知识成果。系统"
        "重点解决课程资料查找、网页正文获取、知识点整理、教学资源归档和后续检索过程中"
        "效率低、格式不统一、来源不易追溯的问题。"
    ),
    B(
        "提供从关键词检索到公开网页资源采集、正文清洗、课程知识抽取、成果保存的连续处理链路。",
        "支持用户指定知识类型，或由系统根据教学资源正文自动识别适合的知识类型。",
        "保存页面来源信息、正文副本、结构化 JSON 成果和处理追踪信息，便于课程资料核对。",
        "提供去重精炼和语义检索能力，使积累后的课程知识成果可以继续整理、查询和复用。",
        "为教学平台或教学 Agent 提供知识库支撑，服务于课程资料管理、知识问答和 RAG 检索增强。",
    ),
    H(2, "1.4 适用范围"),
    P(
        "本软件适用于高校课程资源整理、软件工程课程知识沉淀、教学平台资料归档、"
        "公开教学网页整理、课程知识库构建、面向教学 Agent 的知识数据准备以及 RAG 检索增强等场景。"
        "系统当前重点支持软件工程、数据流图、规则、示例、模板、理论知识等类型的课程资料抽取，"
        "也可扩展到其他教学资源和结构化知识类型。"
    ),
    H(2, "1.5 基本功能概述"),
    T(
        ("功能", "说明"),
        (
            ("公开教学资源采集", "通过关键词检索或手动 URL 输入获取公开网页正文，并保存解析后的页面数据。"),
            ("课程知识抽取", "根据文本内容生成结构化 JSON 成果，支持自动识别或手动指定知识类型。"),
            ("质量控制", "对抽取结果执行格式、字段和内容校验，不合格结果不作为有效成果落盘。"),
            ("成果组织", "按日期、时间段、域名和页面标识组织成果目录，保留正文、元数据和追踪文件。"),
            ("去重精炼", "对已保存成果进行相似内容识别、合并和统计，减少知识库冗余。"),
            ("语义检索", "在已构建向量索引后，根据自然语言查询返回相关课程知识片段。"),
        ),
        (3.6, 12.2),
    ),
    H(1, "第二章 运行环境与配置"),
    H(2, "2.1 运行环境"),
    T(
        ("类别", "要求或组件", "说明"),
        (
            ("操作系统", "Windows、Linux、macOS", "支持常见桌面或服务器环境。"),
            ("Python", "Python >=3.13", "项目配置中声明的运行版本要求。"),
            ("Web 框架", "Flask、Flask-CORS", "用于提供浏览器控制台和后端接口。"),
            ("网络请求", "HTTPX、Requests", "用于搜索接口请求和公开教学网页访问。"),
            ("网页解析", "BeautifulSoup、HTML 清洗组件", "用于提取正文、标题、链接和可读文本。"),
            ("动态渲染", "Playwright", "用于常规正文获取失败或正文过短时处理复杂页面。"),
            ("检索服务", "SerpAPI", "用于关键词搜索并返回候选网页地址。"),
            ("智能抽取", "兼容 OpenAI 格式的 LLM API", "用于知识类型识别、结构化抽取和精炼。"),
            ("向量检索", "ChromaDB 或同类组件", "可选，用于语义索引构建和自然语言检索。"),
            ("数据格式", "JSON、Markdown、纯文本", "用于保存结构化成果、正文副本和追踪信息。"),
        ),
        (3.0, 4.3, 8.5),
    ),
    H(2, "2.2 配置文件与密钥"),
    P(
        "本软件通过环境变量和配置文件读取外部服务参数。用户需根据部署环境配置检索服务密钥、"
        "智能抽取服务地址、智能抽取服务密钥、模型名称、向量检索配置以及可选代理配置。"
        "密钥仅用于本机或服务器运行时访问外部服务，不写入公开文档和知识成果正文。"
    ),
    B(
        "检索服务配置用于执行关键词搜索，缺失时关键词搜索功能不可用，但手动 URL 采集仍可使用。",
        "智能抽取服务配置用于类型识别、知识抽取和精炼，缺失时只能完成公开网页资源采集和正文保存。",
        "向量检索配置用于语义检索，未构建索引时语义检索页会返回初始化或查询错误。",
        "网络代理配置为可选增强能力，用于提高特定网络环境下的网页访问稳定性。",
    ),
    H(2, "2.3 启动与访问方式"),
    P(
        "本软件提供 Web 控制台入口。用户进入软件目录后执行启动脚本，命令行会输出本地访问地址。"
        "启动脚本默认访问地址为 http://localhost:8000；若端口被占用，Web 应用会尝试切换到"
        "备用端口。代码中保留的历史打印信息不影响实际服务端口，以运行输出为准。"
    ),
    T(
        ("入口", "说明"),
        (
            ("python start_web.py", "推荐入口，启动 Web 控制台并提示访问地址。"),
            ("web/app.py", "Web 应用主体，默认监听 8000，端口占用时尝试 8001。"),
            ("server.py", "扩展服务入口，可提供教学资源采集和对外工具调用能力。"),
        ),
        (4.0, 11.8),
    ),
    H(1, "第三章 系统总体架构"),
    H(2, "3.1 架构概述"),
    P(
        "本软件采用 Web 控制台加后台处理流水线的结构。浏览器端负责输入关键词、URL、知识类型、"
        "采集数量、精炼范围和检索条件；后端负责创建任务、调用外部检索服务、执行公开网页资源采集、"
        "清洗正文、调用知识抽取流水线、写入成果文件并向前端返回进度和结果。"
    ),
    C(
        "用户浏览器\n"
        "  -> Web 控制台\n"
        "  -> 任务管理\n"
        "  -> 检索与 URL 队列\n"
        "  -> 公开网页资源采集与正文清洗\n"
        "  -> 知识类型识别\n"
        "  -> 专业知识抽取\n"
        "  -> 质量校验\n"
        "  -> 成果写入\n"
        "  -> 去重精炼与语义检索"
    ),
    H(2, "3.2 分层结构"),
    T(
        ("层次", "组成", "职责"),
        (
            ("表现层", "HTML、CSS、JavaScript", "提供采集、精炼、语义检索页签和结果展示区域。"),
            ("接口层", "Flask 路由", "接收用户请求，返回任务状态、日志、结果和配置列表。"),
            ("任务层", "任务管理器、后台线程", "控制单次任务生命周期，维护进度、日志和执行状态。"),
            ("采集层", "搜索服务、资源采集器、HTML 清洗器", "获取候选 URL，获取页面正文并生成可抽取文本。"),
            ("抽取层", "类型路由、专业抽取器、校验器", "识别知识类型，生成结构化 JSON 并校验质量。"),
            ("存储层", "artifacts 写入器、URL 索引、向量索引", "组织保存成果，维护去重信息和检索索引。"),
        ),
        (2.8, 4.8, 8.2),
    ),
    H(2, "3.3 主要数据对象"),
    P(
        "系统处理过程中主要流转四类数据对象：用户输入、网页解析结果、知识抽取结果和成果存储记录。"
        "用户输入包括关键词、采集数量、手动 URL、知识类型和检索意图；网页解析结果包括标题、"
        "来源地址、清洗正文和结构化段落；知识抽取结果以 JSON 形式保存；成果存储记录用于关联"
        "来源、正文、抽取结果、追踪过程和统计指标。"
    ),
    H(1, "第四章 功能模块设计"),
    H(2, "4.1 模块职责总表"),
    T(
        ("模块", "输入", "输出", "设计说明"),
        (
            ("Web 控制台", "用户操作", "接口请求", "集中提供采集、精炼、语义检索三类操作入口。"),
            ("任务管理", "接口请求", "状态、进度、日志", "限制同类任务并发，避免日志交错和资源冲突。"),
            ("检索模块", "关键词、数量", "URL 列表", "调用检索服务获取候选网页地址。"),
            ("URL 索引", "URL", "是否重复", "记录已采集地址，批量采集时跳过重复 URL。"),
            ("公开网页资源采集", "URL", "页面解析结果", "常规访问优先，必要时使用动态渲染能力。"),
            ("正文清洗", "HTML 或页面数据", "清洗正文", "去除导航、广告、脚本等干扰内容。"),
            ("类型识别", "标题、正文", "知识类型列表", "按可用类型判断适合的抽取方向。"),
            ("专业抽取", "正文、类型", "结构化 JSON", "对不同知识类型调用对应抽取器。"),
            ("质量校验", "抽取结果、原文", "通过或失败", "校验字段完整性、内容相关性和基本格式。"),
            ("成果写入", "页面数据、抽取结果", "成果目录", "写入元数据、正文、追踪和指标文件。"),
            ("去重精炼", "历史成果", "精炼统计和合并结果", "扫描 artifacts，识别重复并生成整理结果。"),
            ("语义检索", "自然语言查询", "命中结果列表", "基于向量索引返回相关知识片段。"),
        ),
        (2.9, 3.6, 3.5, 5.8),
    ),
    H(2, "4.2 Web 控制台模块"),
    P(
        "Web 控制台由浏览器页面和 Flask 接口共同组成。页面包含采集、去重精炼和语义检索三个页签。"
        "采集页签提供关键词输入框、按钮“开始构建”、手动 URL 输入框、按钮“手动爬取”、"
        "采集数量下拉框、知识类型下拉框、进度条、执行日志和抽取结果列表。"
    ),
    P(
        "前端通过异步接口发起任务，再定时获取状态。后端返回任务状态、进度百分比、最近日志和"
        "结果摘要，使用户无需直接查看命令行即可掌握任务执行情况。"
    ),
    H(2, "4.3 检索与 URL 管理模块"),
    P(
        "检索模块接收用户输入的关键词和数量上限，调用 SerpAPI 搜索接口获取候选结果。"
        "每条候选结果中提取链接并加入待采集队列。URL 管理模块对地址进行标准化和哈希，"
        "通过本地索引判断该地址是否已经采集过。批量任务中重复 URL 会被跳过，手动 URL 任务"
        "即使发现重复也允许用户重新采集。"
    ),
    H(2, "4.4 公开网页资源采集与正文清洗模块"),
    P(
        "公开网页资源采集模块负责访问 URL、获取页面正文并生成解析结果。系统优先采用普通网络请求方式"
        "处理静态页面；当正文过短、页面结构复杂或需要动态渲染时，可使用 Playwright 等动态能力"
        "增强采集效果。正文清洗模块从网页中提取标题、正文和结构化段落，减少导航、广告、"
        "推荐链接、脚本样式等非正文内容对抽取过程的干扰。"
    ),
    H(2, "4.5 知识类型识别与专业抽取模块"),
    P(
        "知识抽取分为类型识别和专业抽取两个阶段。类型识别阶段根据标题、正文摘要和已启用的"
        "知识类型列表判断适合的类型；如果用户在界面中手动指定类型，系统直接按指定类型处理。"
        "专业抽取阶段根据类型调用对应抽取器，生成具有字段结构的 JSON 知识成果。"
    ),
    B(
        "DFD 概念类用于抽取外部实体、处理过程、数据流、数据存储等概念说明。",
        "DFD 规则类用于抽取命名规范、连接规则、平衡原则、分解规则等约束内容。",
        "DFD 示例类用于抽取业务场景、参与对象、数据流向和示例说明。",
        "软件工程理论类用于抽取需求分析、结构化分析、系统工程等理论知识。",
    ),
    H(2, "4.6 质量校验与成果写入模块"),
    P(
        "质量校验模块对抽取结果进行基本检查，包括结构是否符合预期、必要字段是否存在、内容是否"
        "与原文相关、是否存在明显空结果等。通过校验的结果交由成果写入模块保存；未通过校验的"
        "结果记录错误信息，不作为有效知识成果输出。"
    ),
    H(2, "4.7 去重精炼与语义检索模块"),
    P(
        "去重精炼模块扫描历史成果目录，按日期、时间段和知识类型聚合待处理文件，识别相近或重复"
        "内容并执行合并、跳过或保留操作。语义检索模块依赖已构建的向量索引，将用户输入的问题"
        "转换为查询条件，在知识集合中返回相关片段、相似度分数、来源和类型信息。"
    ),
    H(1, "第五章 数据流程设计"),
    H(2, "5.1 关键词采集流程"),
    C(
        "输入关键词和采集数量\n"
        "  -> 检查任务状态\n"
        "  -> 调用检索服务\n"
        "  -> 提取 URL 列表\n"
        "  -> URL 去重\n"
        "  -> 逐个采集网页\n"
        "  -> 清洗正文并保存解析结果\n"
        "  -> 类型识别或使用指定类型\n"
        "  -> 专业抽取\n"
        "  -> 质量校验\n"
        "  -> 写入 artifacts 成果目录\n"
        "  -> 返回日志和结果摘要"
    ),
    H(2, "5.2 手动 URL 采集流程"),
    P(
        "手动 URL 流程跳过关键词检索，直接以用户输入的网址作为采集对象。系统先校验 URL 是否为空，"
        "再检查当前是否已有任务运行。若地址曾被采集，系统在日志中提示，但仍允许用户重新采集。"
        "后续的公开网页资源采集、正文清洗、知识抽取、质量校验和成果写入流程与关键词采集流程一致。"
    ),
    H(2, "5.3 去重精炼流程"),
    C(
        "选择日期和时间段\n"
        "  -> 扫描 se_kb/artifacts\n"
        "  -> 过滤非知识成果文件\n"
        "  -> 按知识类型分组\n"
        "  -> 计算相似度或重复关系\n"
        "  -> 生成预览报告\n"
        "  -> 正式执行合并或跳过\n"
        "  -> 输出精炼日志和统计信息"
    ),
    H(2, "5.4 语义检索流程"),
    P(
        "语义检索流程以用户输入的问题、检索意图和返回数量为输入。系统检查向量检索组件是否已初始化，"
        "再向相应集合发起查询。返回结果包含文本片段、相似度分数、来源、知识类型和集合名称，"
        "前端以列表方式展示给用户。"
    ),
    H(1, "第六章 数据组织与存储设计"),
    H(2, "6.1 目录组织"),
    T(
        ("目录或文件", "用途"),
        (
            ("config/", "保存系统提示、模板、解析器和路径配置。"),
            ("data/url_index.json", "保存已采集 URL 的索引和采集时间。"),
            ("data/raw/", "保存历史原始响应或中间采集结果。"),
            ("data/parsed/", "保存解析后的结构化页面数据。"),
            ("se_kb/artifacts/", "保存按任务批次组织的知识抽取成果。"),
            ("se_kb/diagrams/、se_kb/theory/ 等", "保存经整理后的领域知识库内容。"),
            ("web/templates/、web/static/", "保存 Web 控制台页面和前端资源。"),
            ("logs/", "保存运行日志或调试日志。"),
        ),
        (5.0, 10.8),
    ),
    H(2, "6.2 成果目录结构"),
    C(
        "se_kb/artifacts/\n"
        "  YYYY/\n"
        "    MM/\n"
        "      DD/\n"
        "        HH_MM/\n"
        "          domain.example/\n"
        "            title-slug-hash/\n"
        "              metadata.json\n"
        "              parsed.json\n"
        "              document.md\n"
        "              trace.json\n"
        "              metrics.json\n"
        "              errors.json\n"
        "              <knowledge_type>_<slug>.json"
    ),
    H(2, "6.3 成果文件说明"),
    T(
        ("文件", "内容", "作用"),
        (
            ("metadata.json", "来源 URL、标题、知识类型等元信息", "用于溯源和检索。"),
            ("parsed.json", "网页标题、正文、结构化段落等解析结果", "用于复核采集质量。"),
            ("document.md", "清洗后的正文文本副本", "用于人工阅读和对照原文。"),
            ("trace.json", "类型识别、抽取、校验等过程追踪", "用于排查抽取过程。"),
            ("metrics.json", "令牌、耗时或统计指标", "用于评估资源消耗。"),
            ("errors.json", "失败或异常信息", "用于定位处理失败原因。"),
            ("知识类型 JSON", "最终结构化知识成果", "用于知识库导入、精炼和检索。"),
        ),
        (3.5, 6.2, 6.1),
    ),
    H(2, "6.4 URL 索引设计"),
    P(
        "URL 索引使用本地 JSON 文件保存。系统对 URL 进行小写化和去除尾部斜杠等标准化处理，"
        "再生成短哈希作为索引键。索引值包含原始 URL、标题和采集时间。该设计避免重复网页在批量"
        "任务中被反复处理，同时保留必要的来源信息。"
    ),
    H(1, "第七章 关键处理流程设计"),
    H(2, "7.1 任务状态控制"),
    P(
        "系统在后端维护任务状态、进度、日志和结果列表。用户发起任务时，后端先检查当前状态。"
        "若已有采集或精炼任务正在运行，接口返回提示信息；若状态允许，则创建后台线程执行任务。"
        "该机制避免多个长任务同时写入日志和成果目录，降低资源竞争风险。"
    ),
    H(2, "7.2 公开网页资源采集处理链路"),
    P(
        "公开网页资源采集处理链路由 URL 输入、网络访问、正文解析、清洗、保存组成。资源采集器获取网页正文后，"
        "提取标题、正文、来源地址和段落结构，并保存为解析文件。若采集失败，系统在日志中记录失败"
        "原因并继续处理其他 URL，避免单个网页失败导致整批任务中断。"
    ),
    H(2, "7.3 知识抽取处理链路"),
    P(
        "知识抽取处理链路以解析后的页面数据为输入。系统先准备标题、来源地址和正文摘要，"
        "再根据用户指定类型或自动识别结果选择抽取器。抽取器调用智能抽取服务生成 JSON，"
        "随后质量校验模块对结果进行检查。通过校验后，成果写入器按照日期、域名和标题哈希"
        "组织目录并写入相关文件。"
    ),
    H(2, "7.4 精炼处理链路"),
    P(
        "精炼处理链路从 artifacts 目录读取历史成果，排除 metadata、parsed、trace、metrics、"
        "errors 等辅助文件，只处理实际知识成果文件。系统按知识类型分组后进行相似度检测，"
        "给出预览统计。正式精炼时，系统根据检测结果合并相近内容或跳过重复内容，并生成处理日志。"
    ),
    H(2, "7.5 检索处理链路"),
    P(
        "语义检索处理链路依赖向量索引。索引构建程序将 JSON 知识成果切分为文本片段，写入向量"
        "存储集合。查询时，检索器根据用户输入和检索意图选择集合或过滤条件，返回相似度最高的"
        "若干条记录。前端显示内容片段、来源、类型和相似度，便于用户继续核对。"
    ),
    H(1, "第八章 异常处理与运行保障"),
    H(2, "8.1 常见异常处理"),
    T(
        ("异常场景", "系统处理方式", "用户处理建议"),
        (
            ("关键词为空", "接口返回错误提示，不启动任务。", "补充搜索关键词后重新提交。"),
            ("已有任务运行", "拒绝新任务并返回已有任务提示。", "等待当前任务完成后再操作。"),
            ("检索密钥缺失", "日志提示未配置检索密钥，返回空 URL 列表。", "补充检索服务配置或改用手动 URL。"),
            ("搜索请求失败", "重试后仍失败则记录错误并结束任务。", "检查网络、密钥和服务额度。"),
            ("网页正文获取失败", "记录失败日志，批量任务继续处理其他 URL。", "更换网页或检查网络访问能力。"),
            ("解析文件缺失", "标记任务错误，不继续抽取。", "检查资源采集器输出目录和写入权限。"),
            ("抽取无结果", "日志提示未抽取到知识，任务可正常结束。", "调整知识类型或更换内容更明确的网页。"),
            ("校验不通过", "不保存为有效成果，记录失败原因。", "查看 trace 或 errors 文件定位问题。"),
            ("语义检索未初始化", "接口返回初始化错误。", "先构建向量索引或重载知识库连接。"),
        ),
        (3.6, 6.3, 6.0),
    ),
    H(2, "8.2 日志与追踪"),
    P(
        "系统在前端展示最近执行日志，并在成果目录中保存 trace 和 metrics 等追踪文件。"
        "前端日志用于用户实时观察任务状态，追踪文件用于开发或维护人员复盘类型识别、"
        "抽取调用、校验结果和资源使用情况。"
    ),
    H(2, "8.3 数据保护"),
    P(
        "系统不会将外部服务密钥写入知识成果文件。用户应将配置文件和密钥文件保存在受控环境中，"
        "迁移或备份知识成果时应区分成果目录和密钥配置。对外分享成果时，建议先检查 metadata、"
        "trace 和日志文件中是否包含不适合公开的来源信息或运行环境信息。"
    ),
    H(1, "第九章 使用说明"),
    H(2, "9.1 启动软件"),
    B(
        "安装 Python 和依赖库，并根据部署说明完成必要配置。",
        "进入软件目录，执行 python start_web.py。",
        "在命令行输出中查看访问地址，通常为 http://localhost:8000。",
        "用浏览器打开访问地址，确认页面加载成功并显示采集、去重精炼、语义检索功能入口。",
    ),
    H(2, "9.2 关键词采集"),
    B(
        "在采集页签输入关键词，例如“数据流图 绘制步骤 教程”。",
        "选择采集数量和知识类型；界面中“采取知识”对应本文所称“教学资源采集”。",
        "点击“开始构建”后查看进度条和执行日志。",
        "任务结束后，在抽取结果列表中查看标题、知识类型和抽取状态。",
        "如需核对详细内容，可到 artifacts 成果目录查看 JSON、正文和追踪文件。",
    ),
    H(2, "9.3 手动 URL 采集"),
    B(
        "复制目标网页的完整 URL。",
        "粘贴到手动 URL 输入框。",
        "点击“手动爬取”。",
        "等待日志显示采集、抽取和保存过程。",
        "任务完成后查看抽取结果和成果目录。",
    ),
    H(2, "9.4 去重精炼"),
    B(
        "打开去重精炼页签。",
        "选择日期和时间段，或保留全部范围。",
        "需要先查看影响范围时，使用预览分析。",
        "确认后点击开始精炼。",
        "查看精炼日志和统计结果，确认合并、跳过和错误数量。",
    ),
    H(2, "9.5 语义检索"),
    B(
        "确认向量索引已经构建完成。",
        "打开语义检索页签，输入自然语言问题。",
        "选择检索意图和结果数量。",
        "点击开始检索，查看返回的文本片段、来源、类型和相似度信息。",
    ),
    H(1, "第十章 维护、安全与版本管理"),
    H(2, "10.1 维护方式"),
    P(
        "维护人员可通过配置文件调整检索服务、智能抽取服务、模型名称、代理参数和向量检索参数。"
        "系统提示模板、知识类型定义和抽取器可按知识领域扩展。新增知识类型时，应同步补充类型注册、"
        "抽取逻辑、校验规则和必要的前端选项。"
    ),
    H(2, "10.2 备份与迁移"),
    P(
        "知识成果主要保存在 se_kb/artifacts 及相关知识库目录中。迁移软件时，应同时复制成果目录、"
        "必要配置和向量索引目录。涉及密钥的配置文件应通过安全渠道迁移，不应随公开成果包一起分发。"
    ),
    H(2, "10.3 版本管理"),
    P(
        f"本文档对应的软件版本为{VERSION}。后续版本若新增功能模块、调整界面名称、改变成果目录结构"
        "或更换主要技术依赖，应同步更新软件设计说明书、用户操作说明和程序鉴别材料中的名称、"
        "版本、模块说明及运行环境描述。"
    ),
    H(1, "附录 A 术语说明"),
    T(
        ("术语", "含义"),
        (
            ("教学资源采集", "面向公开教学网页访问页面并提取标题、正文和结构化页面信息的过程。"),
            ("正文清洗", "去除导航、广告、脚本样式等干扰内容，保留可读正文的过程。"),
            ("知识抽取", "将正文转换为具有字段结构的 JSON 知识成果的过程。"),
            ("知识类型", "抽取结果所属类别，例如 DFD 概念、DFD 规则、DFD 示例、软件工程理论等。"),
            ("成果目录", "按日期、时间、域名和页面标识保存知识成果的目录结构。"),
            ("去重精炼", "识别重复或相近知识内容，并进行合并、跳过或整理的过程。"),
            ("语义检索", "基于向量索引和自然语言查询查找相关知识片段的过程。"),
            ("LLM API", "用于类型识别、知识抽取和精炼的大语言模型服务接口。"),
        ),
        (3.5, 12.3),
    ),
    H(1, "附录 B 登记信息一致性说明"),
    T(
        ("项目", "正式填写口径"),
        (
            ("软件名称", SOFTWARE_NAME),
            ("版本号", VERSION),
            ("文档名称", DOC_TITLE),
            ("界面显示名称", "首次说明其与正式软件名称对应，正文不作为申报名称使用。"),
            ("著作权人", f"{COPYRIGHT_OWNER_FORMAL}。"),
            ("提交材料", "正式 Word 或 PDF 文档应保持页码连续、页眉页脚一致、章节内容连贯。"),
        ),
        (4.2, 11.6),
    ),
    P("全文结束。"),
]


def _insert_before_text(target_text: str, additions: list[Element]) -> None:
    for index, element in enumerate(ELEMENTS):
        if element.text == target_text:
            ELEMENTS[index:index] = additions
            return
    raise ValueError(f"Insert target not found: {target_text}")


MODULE_EXPANSION: list[Element] = [
    H(2, "4.8 Web 控制台功能细化"),
    P(
        "Web 控制台面向最终用户提供集中操作入口。用户在页面上完成关键词输入、手动 URL 输入、"
        "知识类型选择、采集数量选择、精炼范围选择和语义检索查询。控制台不直接暴露内部文件路径、"
        "模型调用细节和密钥配置，主要呈现任务状态、进度百分比、执行日志和结果摘要。"
    ),
    P(
        "该模块的输入包括用户在页面组件中填写的查询词、URL、下拉选项和按钮动作。输出包括后端"
        "接口请求、页面状态变化、日志列表刷新和结果区域渲染。处理结果表现为页面上的等待中、"
        "运行中、完成、错误等状态，以及抽取结果、精炼统计或检索结果列表。"
    ),
    H(2, "4.9 检索与链接获取功能细化"),
    P(
        "检索与链接获取模块负责把用户输入的主题词转换为候选网页地址。用户触发“开始构建”后，"
        "系统读取关键词和采集数量，调用检索服务获得若干条搜索结果。模块只把可打开的 URL 交给"
        "后续采集链路，检索结果标题主要用于日志提示和人工识别。"
    ),
    P(
        "该模块的输入为关键词、数量上限和检索服务配置；输出为 URL 列表和日志信息。页面表现为"
        "执行日志中出现“搜索”“找到结果”“跳过重复 URL”等提示。若检索服务不可用，系统不会继续"
        "执行批量采集，而是将失败原因反馈给用户，避免产生来源不明的空结果。"
    ),
    H(2, "4.10 公开网页资源采集功能细化"),
    P(
        "公开网页资源采集模块以 URL 为基本处理单元。对每个 URL，系统先检查是否已存在于 URL 索引中，"
        "再调用资源采集器访问网页。普通网页优先采用轻量请求方式获取正文，复杂页面可通过动态渲染方式"
        "补充处理。采集结果会被整理为标题、来源地址、正文、段落结构和页面解析文件。"
    ),
    P(
        "该模块的处理结果一方面写入 data 或 artifacts 相关目录，另一方面通过日志向页面反馈。"
        "用户看到的是当前正在处理第几个网页、是否获取成功、是否跳过重复网页等信息。单个网页失败"
        "不会阻塞整批任务，后续网页仍会继续处理。"
    ),
    H(2, "4.11 正文清洗功能细化"),
    P(
        "正文清洗模块负责将网页原始内容转换为适合知识抽取的文本。系统会过滤页面中的导航菜单、"
        "广告、脚本样式、推荐链接和无关页脚，保留正文标题、段落、小标题和必要的上下文信息。"
        "清洗后的正文既用于模型抽取，也会作为 document.md 或解析文件内容保存，方便人工核对。"
    ),
    P(
        "该模块的输入为网页 HTML、资源采集器解析结果或结构化段落；输出为 clean_text、markdown 或"
        "sections 等可读文本结构。处理结果要求语义完整、噪声较少、来源可追溯，以减少后续模型"
        "误判和无效抽取。"
    ),
    H(2, "4.12 知识抽取功能细化"),
    P(
        "知识抽取模块根据知识类型生成结构化成果。自动识别模式下，系统先判断正文更接近哪类知识；"
        "手动指定模式下，系统按用户指定类型调用抽取器。不同类型对应不同的字段组织方式，例如概念"
        "类侧重定义与要素，规则类侧重约束和适用条件，示例类侧重业务场景和数据流向。"
    ),
    P(
        "该模块的输入为页面标题、来源地址、清洗正文和知识类型；输出为 JSON 格式知识成果、抽取"
        "追踪信息和资源统计信息。页面表现为抽取结果列表，列表通常包含标题、类型和处理状态。"
    ),
    H(2, "4.13 成果管理功能细化"),
    P(
        "成果管理模块把一次任务的各类产物组织为可追溯目录。每个网页成果目录中既保存结构化知识"
        "JSON，也保存原始解析数据、正文副本、元数据、处理追踪和指标信息。该设计使用户后续可以"
        "回看来源网页、对照正文、检查抽取字段，也便于将成果导入向量索引或其他知识库系统。"
    ),
    P(
        "该模块的输入为网页解析结果、知识抽取结果、追踪信息、指标信息和错误信息；输出为按日期、"
        "时间、域名、页面标识分层的文件夹。若目标目录不存在，系统自动创建；若写入失败，任务日志"
        "应记录错误，便于检查磁盘权限或路径配置。"
    ),
    H(2, "4.14 精炼与检索功能细化"),
    P(
        "精炼功能用于处理已经积累的成果，重点解决重复内容和相近内容过多的问题。用户选择日期和"
        "时间段后，系统扫描对应 artifacts 目录，按知识类型分组后识别重复或相似成果。预览模式只"
        "展示分析结论，正式执行才会写入整理结果或更新统计。"
    ),
    P(
        "语义检索功能用于在结构化知识中快速定位内容。用户输入自然语言问题后，系统根据检索意图"
        "选择集合或过滤条件，返回相关片段、相似度分数、来源和类型。该功能依赖向量索引，未建索引"
        "时不会返回有效命中结果。"
    ),
]


FLOW_EXPANSION: list[Element] = [
    H(2, "5.5 批量采集流程细化"),
    P(
        "用户触发条件为在采集页签输入关键词并点击“开始构建”。系统首先检查关键词是否为空、任务"
        "是否正在运行、采集数量是否合法。检查通过后，系统进入检索阶段，得到候选 URL；随后执行"
        "URL 去重，将已采集地址放入跳过列表，将新地址放入待处理列表。"
    ),
    P(
        "对待处理 URL，系统逐个执行公开网页资源采集和正文清洗。每个 URL 的处理结果会形成解析文件路径，"
        "并追加到后续抽取队列。全部采集完成后，系统进入知识抽取阶段，对每个解析文件执行类型识别、"
        "专业抽取和质量校验。最终输出包括前端结果摘要、URL 索引更新记录和 artifacts 成果目录。"
    ),
    H(2, "5.6 单页采集流程细化"),
    P(
        "用户触发条件为在手动 URL 输入框中填写完整网页地址并点击“手动爬取”。系统首先判断 URL"
        "是否为空；若为空，接口直接返回提示。若当前已有主任务运行，系统拒绝启动新任务。通过检查"
        "后，即使 URL 已存在索引，也允许重新采集，以满足用户需要复核或刷新页面内容的场景。"
    ),
    P(
        "单页采集的输出结果与批量采集一致，包括解析文件、正文副本、知识 JSON、追踪文件和页面"
        "结果摘要。与批量采集不同的是，单页流程不需要检索服务，也不会因为检索密钥缺失而不可用。"
    ),
    H(2, "5.7 抽取校验流程细化"),
    P(
        "抽取校验流程在模型返回结构化 JSON 后执行。系统读取抽取结果，检查类型是否匹配、字段是否"
        "符合预期、必要内容是否为空、结果是否能与正文建立基本关联。通过校验的成果写入对应知识类型"
        "文件；未通过校验的成果不会作为有效结果保存。"
    ),
    P(
        "该流程的输出包括成功成果、失败原因和追踪信息。追踪信息不仅用于调试，也用于解释一次任务"
        "为什么只产生部分结果。例如，某些网页虽然正文获取成功，但正文与目标知识类型不匹配，系统会记录"
        "抽取为空或校验不通过，而不是生成质量较差的知识条目。"
    ),
    H(2, "5.8 索引构建与查询流程细化"),
    P(
        "语义检索依赖索引构建。构建索引时，系统扫描正式知识库或 artifacts 中的 JSON 文件，排除"
        "元数据、追踪、指标和错误文件，将实际知识内容拆分为适合检索的文本片段。每个片段携带来源、"
        "知识类型、集合名称和字段信息，写入向量存储。"
    ),
    P(
        "查询时，系统根据用户输入和检索意图选择查询范围，将自然语言问题与向量集合进行匹配，并按"
        "相似度排序返回结果。输出结果不直接修改知识库，只作为阅读和定位依据。"
    ),
]


EXCEPTION_EXPANSION: list[Element] = [
    H(2, "8.4 任务并发异常说明"),
    P(
        "当用户连续点击“开始构建”“手动爬取”或“开始精炼”时，系统可能收到多个任务请求。为避免"
        "多个任务同时写入状态、日志和成果目录，任务管理器会检查当前状态。若已有任务处于 running，"
        "新请求被拒绝，并在页面上提示已有任务正在运行。"
    ),
    P(
        "这种处理方式可以保证日志顺序清晰、结果目录不会互相覆盖，也减少了教学资源采集和模型调用同时"
        "占用资源造成的失败风险。用户只需要等待进度完成后再次发起任务。"
    ),
    H(2, "8.5 网络访问异常说明"),
    P(
        "网络访问异常可能发生在检索服务调用、网页正文获取、动态渲染和智能抽取服务请求等阶段。系统会"
        "在日志中记录失败位置和错误摘要。对批量教学资源采集而言，单个 URL 的网络失败不会导致整批任务"
        "立即终止，系统会继续尝试后续 URL。"
    ),
    P(
        "如果检索服务整体不可用，则批量任务无法获得 URL 列表，任务会以错误状态结束。此时用户可以"
        "检查网络和密钥，也可以改用手动 URL 采集，以绕过检索阶段继续处理已知网页。"
    ),
    H(2, "8.6 抽取服务异常说明"),
    P(
        "智能抽取服务异常包括 API 地址配置错误、密钥无效、模型无响应、返回内容无法解析为 JSON、"
        "抽取内容为空等情况。系统会把错误写入日志或 trace 信息中，并将该网页标记为未产生有效抽取"
        "结果。已经成功采集的正文副本仍会保留，便于用户稍后修复配置后重新抽取。"
    ),
    P(
        "系统不把异常返回内容强行写成知识成果。这样可以避免把不完整、格式错误或与正文无关的内容"
        "混入正式知识库。"
    ),
    H(2, "8.7 存储异常说明"),
    P(
        "存储异常通常与目录权限、磁盘空间、文件名过长或路径配置有关。成果写入器在写入 metadata、"
        "parsed、document、trace、metrics 和知识 JSON 时，需要目标目录具备写入权限。若写入失败，"
        "系统应记录失败文件和异常信息。"
    ),
    P(
        "用户迁移软件目录或切换运行账户后，应重点检查 data、se_kb/artifacts、向量索引目录和日志"
        "目录的写入权限，确保采集、抽取、精炼和检索索引构建均可正常落盘。"
    ),
]


USAGE_EXPANSION: list[Element] = [
    H(2, "9.6 典型场景一：课程资料整理"),
    P(
        "用户在学习软件工程课程时，可围绕“数据流图”“需求分析”“结构化设计”等主题进行关键词采集。"
        "首次使用时建议选择较小的采集数量，例如 3 个网页，观察日志和抽取结果是否符合预期。确认"
        "配置正常后，再扩大采集数量，逐步形成课程知识资料集。"
    ),
    P(
        "该场景的输出重点是概念、规则和理论类知识。用户可以通过 artifacts 中的正文副本核对来源，"
        "也可以在后续构建索引后用自然语言查询课程相关内容。"
    ),
    H(2, "9.7 典型场景二：教学平台资源沉淀"),
    P(
        "课程团队可以把教学大纲、公开课程资料、教学平台说明页或公开教程链接通过手动 URL 方式逐篇采集。"
        "手动 URL 方式适合来源明确、质量较高、需要保留出处的教学资料。系统会为每篇资料保存元数据和正文副本，便于"
        "教师或助教复核和后续归档。"
    ),
    P(
        "当同一课程主题积累较多成果后，教学平台维护人员可使用去重精炼功能识别重复内容。正式精炼前建议先执行预览，"
        "查看重复数量、知识类型分布和预计处理情况，再决定是否写入整理结果。"
    ),
    H(2, "9.8 典型场景三：知识库检索复用"),
    P(
        "当课程知识成果积累到一定规模后，用户可运行索引构建程序，将 JSON 成果转换为向量检索集合。"
        "之后即可在语义检索页输入问题，例如“DFD 外部实体如何命名”或“父子图平衡原则是什么”，"
        "系统会返回相关课程知识片段和来源信息。"
    ),
    P(
        "该场景强调复用和定位。用户不需要记住具体文件夹位置，只需通过自然语言查询找到相关片段，"
        "再根据来源字段回到成果目录核对完整 JSON 或正文副本。"
    ),
    H(2, "9.9 软件真实运行界面与完整使用流程"),
    P(
        "本节结合软件真实运行界面展示从系统进入、任务配置、执行处理、结果查看、手动 URL 采集、"
        "去重精炼、语义检索到成果文件查看的连贯使用流程。截图内容用于说明软件主要功能和技术特点，"
        "正式提交前应确认画面中不包含账号、密钥、浏览器书签或其他与软件功能无关的敏感信息。"
    ),
    H(3, "9.9.1 打开系统首页"),
    P("用户启动 Web 控制台后，在浏览器中打开系统首页，首页显示采集、去重精炼和语义检索三个主要功能入口。"),
    IMG("output/doc/screenshots/figure-9-1-system-home.png", "图9-1 系统首页与采集功能入口"),
    H(3, "9.9.2 设置关键词采集任务"),
    P("用户在采集页签输入课程相关关键词，选择采集数量和知识类型，为后续公开教学资源采集和结构化抽取提供任务参数。"),
    IMG("output/doc/screenshots/figure-9-2-keyword-settings.png", "图9-2 关键词、采集数量与知识类型设置界面"),
    H(3, "9.9.3 执行采集与抽取任务"),
    P("任务启动后，系统显示运行状态、进度条和执行日志，日志用于连续展示检索、网页正文获取、知识抽取和校验等处理阶段。"),
    IMG("output/doc/screenshots/figure-9-3-task-progress-log.png", "图9-3 任务运行进度与执行日志界面"),
    H(3, "9.9.4 查看抽取结果"),
    P("任务完成后，用户可在抽取结果区域查看知识类型、来源信息和结构化知识摘要，并可进一步到成果目录核对完整文件。"),
    IMG("output/doc/screenshots/figure-9-4-extraction-results.png", "图9-4 抽取结果列表界面"),
    H(3, "9.9.5 手动URL采集"),
    P("当用户已知目标教学网页地址时，可在手动 URL 输入框中填写网页地址，并指定知识类型后执行单页采集与抽取。"),
    IMG("output/doc/screenshots/figure-9-5-manual-url.png", "图9-5 手动URL采集界面"),
    H(3, "9.9.6 去重精炼"),
    P("去重精炼页用于选择成果日期和时间段，支持预览分析和试运行，以便在正式整理前确认处理范围。"),
    IMG("output/doc/screenshots/figure-9-6-refine-settings.png", "图9-6 去重精炼参数设置界面"),
    P("精炼完成后，系统展示处理文件数、生成条目数和去重数量等统计信息，便于用户确认知识库整理效果。"),
    IMG("output/doc/screenshots/figure-9-7-refine-results.png", "图9-7 去重精炼结果统计界面"),
    H(3, "9.9.7 语义检索"),
    P("语义检索页支持用户输入自然语言问题，并选择检索意图和返回数量，用于在已构建知识库中定位相关知识片段。"),
    IMG("output/doc/screenshots/figure-9-8-semantic-input.png", "图9-8 语义检索输入界面"),
    P("检索完成后，系统展示相关知识片段、知识类型、来源集合和相似度信息，可支撑课程知识问答和 RAG 检索增强。"),
    IMG("output/doc/screenshots/figure-9-9-semantic-results.png", "图9-9 语义检索结果界面"),
    H(3, "9.9.8 成果文件查看"),
    P("系统将每次处理产生的元数据、正文副本、解析结果、追踪信息和结构化知识 JSON 保存到成果目录，便于复核、迁移和后续索引构建。"),
    IMG("output/doc/screenshots/figure-9-10-artifacts-files.png", "图9-10 知识成果目录与文件示例"),
]


APPENDIX_EXPANSION: list[Element] = [
    H(1, "附录 C 主要配置项说明"),
    T(
        ("配置项", "用途", "缺失时影响"),
        (
            ("SERPAPI_API_KEY", "用于关键词检索服务认证。", "批量关键词检索不可用，手动 URL 采集不受影响。"),
            ("BASE_URL 或兼容 API 地址", "用于智能抽取服务请求。", "知识类型识别、抽取和精炼可能不可用。"),
            ("OPENAI_API_KEY 或兼容密钥", "用于访问大语言模型服务。", "无法完成模型调用。"),
            ("EMBEDDING_BASE_URL", "用于向量嵌入服务。", "语义索引构建或语义检索不可用。"),
            ("向量库目录", "用于保存 ChromaDB 或同类检索索引。", "检索时无法找到历史索引。"),
            ("网络代理配置", "用于增强特定网络环境下的网页访问能力。", "网络受限时网页正文获取成功率可能下降。"),
        ),
        (4.5, 6.0, 5.3),
    ),
    P(
        "以上配置项根据部署环境填写。正式提交的鉴别材料不应包含真实密钥值，只说明配置项用途。"
        "实际提交前，如需展示运行截图，应确认配置页面、命令行窗口和日志中没有泄露密钥。"
    ),
    H(1, "附录 D 主要成果文件示例说明"),
    T(
        ("成果文件", "示例内容", "核对方式"),
        (
            ("metadata.json", "来源 URL、页面标题、知识类型、任务时间。", "核对成果是否对应目标网页。"),
            ("parsed.json", "网页标题、正文、段落结构、来源地址。", "核对采集和正文清洗是否成功。"),
            ("document.md", "清洗后的正文副本。", "人工阅读并与原网页对照。"),
            ("trace.json", "类型识别、抽取、校验的过程信息。", "排查为什么某类知识成功或失败。"),
            ("metrics.json", "令牌、耗时、统计指标。", "了解任务资源消耗。"),
            ("知识 JSON 文件", "概念、规则、示例、理论等结构化字段。", "作为正式知识成果导入或检索。"),
        ),
        (4.0, 6.4, 5.4),
    ),
    P(
        "成果文件采用统一目录结构保存，使同一网页的来源、正文、抽取结果和追踪信息位于同一目录。"
        "这种组织方式便于迁移、备份、人工复核和后续索引构建。"
    ),
]


_insert_before_text("第五章 数据流程设计", MODULE_EXPANSION)
_insert_before_text("第六章 数据组织与存储设计", FLOW_EXPANSION)
_insert_before_text("第九章 使用说明", EXCEPTION_EXPANSION)
_insert_before_text("第十章 维护、安全与版本管理", USAGE_EXPANSION)
_insert_before_text("全文结束。", APPENDIX_EXPANSION)


INTERNAL_CHECKLIST_TEXT = f"""# {FULL_NAME} 软著提交内部检查清单

本文件仅供提交前内部核对，不属于正式鉴别材料正文。

## 名称与版本
- 申请表、文档、页眉、页脚、源代码材料中的软件名称统一为：{SOFTWARE_NAME}。
- 版本号统一为：{VERSION}。
- 若著作权人名称已确定，可设置环境变量 COPYRIGHT_OWNER 后重新运行生成脚本，使页脚、封面和附录登记信息一次性替换。

## 正式文档
- 正式文档不包含“排版说明”“内部检查”“请替换”等内部提示语。
- 正式文档类型为“软件设计说明书”。
- 界面显示名称只作为界面说明，不作为申报软件名称替代。
- 文档页码连续，A4 版式，页眉页脚清晰。

## 内容核对
- 运行环境包含 Python >=3.13、Flask、HTTPX、BeautifulSoup、Playwright、SerpAPI、LLM API、ChromaDB 或同类向量检索组件。
- 功能描述包含公开教学资源采集、正文清洗、知识类型识别、知识抽取、质量校验、成果写入、去重精炼、语义检索。
- 数据目录描述与当前项目结构一致。
- 启动端口说明以实际运行输出为准。

## 提交材料
- 若文档不足 60 页，提交全部文档。
- 若文档超过 60 页，按登记要求提交前、后各连续 30 页。
- 不为凑页数添加无实质内容。
"""


def iter_headings() -> list[tuple[int, str]]:
    return [(el.level, el.text) for el in ELEMENTS if el.kind == "heading"]


def render_markdown(elements: Iterable[Element]) -> str:
    lines: list[str] = []
    for el in elements:
        if el.kind == "heading":
            lines.append("#" * el.level + " " + el.text)
            lines.append("")
        elif el.kind == "paragraph":
            lines.append(el.text)
            lines.append("")
        elif el.kind == "bullets":
            for item in el.items:
                lines.append(f"- {item}")
            lines.append("")
        elif el.kind == "table":
            lines.append("| " + " | ".join(el.headers) + " |")
            lines.append("| " + " | ".join("---" for _ in el.headers) + " |")
            for row in el.rows:
                lines.append("| " + " | ".join(row) + " |")
            lines.append("")
        elif el.kind == "code":
            lines.append("```text")
            lines.append(el.text)
            lines.append("```")
            lines.append("")
        elif el.kind == "image":
            lines.append(f"![{el.caption}]({el.text})")
            lines.append("")
            lines.append(el.caption)
            lines.append("")
        elif el.kind == "pagebreak":
            lines.append("")
    return "\n".join(lines).strip() + "\n"


def set_run_font(run, name: str = "宋体", size: float = 10.5, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    set_run_font(run, "宋体", 9.5, bold)


def add_page_number(paragraph) -> None:
    paragraph.add_run("第 ")
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_sep)
    run._r.append(text)
    run._r.append(fld_char_end)
    paragraph.add_run(" 页")


def setup_docx_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = document.styles[style_name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(31, 78, 121)


def set_update_fields_on_open(document: Document) -> None:
    settings = document.settings._element
    for existing in settings.findall(qn("w:updateFields")):
        settings.remove(existing)
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_toc_field(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    run._r.append(instr)
    run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def add_docx_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.text = ""
    run = header.add_run(f"{FULL_NAME} 软件设计说明书")
    set_run_font(run, "宋体", 9)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = ""
    run = footer.add_run(f"著作权人：{COPYRIGHT_OWNER}    ")
    set_run_font(run, "宋体", 9)
    add_page_number(footer)
    for run in footer.runs:
        set_run_font(run, "宋体", 9)


def render_docx(elements: Iterable[Element], path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    setup_docx_styles(document)
    add_docx_header_footer(document)
    set_update_fields_on_open(document)

    first_heading = True
    for el in elements:
        if el.kind == "heading":
            if not first_heading and el.level == 1:
                document.add_page_break()
            paragraph = document.add_heading(el.text, level=min(el.level, 3))
            if first_heading:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(12)
                first_heading = False
            else:
                paragraph.paragraph_format.space_before = Pt(8 if el.level == 1 else 4)
                paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs:
                set_run_font(run, "黑体", 17 if el.level == 1 and paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER else 13 if el.level == 1 else 11.5, True)
            if el.text == "目录":
                add_toc_field(document)
        elif el.kind == "paragraph":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Pt(21)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(el.text)
            set_run_font(run, "宋体", 10.5)
        elif el.kind == "bullets":
            for item in el.items:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.left_indent = Pt(18)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.paragraph_format.space_after = Pt(2)
                run = paragraph.add_run(item)
                set_run_font(run, "宋体", 10.3)
        elif el.kind == "table":
            for row in el.rows:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Pt(14)
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.paragraph_format.space_after = Pt(2)
                if len(row) == 2:
                    label = paragraph.add_run(f"{row[0]}：")
                    set_run_font(label, "宋体", 10, True)
                    value = paragraph.add_run(row[1])
                    set_run_font(value, "宋体", 10)
                else:
                    for idx, value_text in enumerate(row):
                        if idx:
                            sep = paragraph.add_run("；")
                            set_run_font(sep, "宋体", 10)
                        label = paragraph.add_run(f"{el.headers[idx]}：")
                        set_run_font(label, "宋体", 10, True)
                        value = paragraph.add_run(value_text)
                        set_run_font(value, "宋体", 10)
        elif el.kind == "code":
            for line in el.text.splitlines():
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Pt(24)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                run = paragraph.add_run(line)
                set_run_font(run, "Consolas", 9.3)
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            document.add_paragraph()
        elif el.kind == "image":
            image_path = resolve_asset(el.text)
            if image_path.exists():
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(3)
                run = paragraph.add_run()
                run.add_picture(str(image_path), width=Cm(15.8))

                caption = document.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.paragraph_format.space_after = Pt(6)
                caption_run = caption.add_run(el.caption)
                set_run_font(caption_run, "宋体", 9.2)
        elif el.kind == "pagebreak":
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def set_fixed_table_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    for child in list(tbl_pr):
        if child.tag in {qn("w:tblW"), qn("w:tblLayout")}:
            tbl_pr.remove(child)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), "9525")
    tbl_pr.append(tbl_w)


def shading(fill: str) -> OxmlElement:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    return shd


def register_pdf_fonts() -> tuple[str, str]:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    heading_font = "STSong-Light"
    simhei = Path("C:/Windows/Fonts/simhei.ttf")
    if simhei.exists():
        pdfmetrics.registerFont(TTFont("SimHei", str(simhei)))
        heading_font = "SimHei"
    return "STSong-Light", heading_font


def html_escape(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\n", "<br/>")
    )


def render_pdf(elements: Iterable[Element], path: Path) -> None:
    body_font, heading_font = register_pdf_fonts()
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="CnBody",
            parent=styles["Normal"],
            fontName=body_font,
            fontSize=9.6,
            leading=13.2,
            firstLineIndent=18,
            alignment=TA_LEFT,
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CnTitle",
            parent=styles["Title"],
            fontName=heading_font,
            fontSize=17,
            leading=22,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#1F4E79"),
            spaceAfter=12,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CnH1",
            parent=styles["Heading1"],
            fontName=heading_font,
            fontSize=13,
            leading=17,
            textColor=colors.HexColor("#1F4E79"),
            spaceBefore=9,
            spaceAfter=5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CnH2",
            parent=styles["Heading2"],
            fontName=heading_font,
            fontSize=11.2,
            leading=15,
            textColor=colors.HexColor("#1F4E79"),
            spaceBefore=6,
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CnBullet",
            parent=styles["Normal"],
            fontName=body_font,
            fontSize=9.4,
            leading=12.8,
            leftIndent=14,
            firstLineIndent=0,
            spaceAfter=2,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CnTable",
            parent=styles["Normal"],
            fontName=body_font,
            fontSize=8.6,
            leading=11.2,
            firstLineIndent=0,
        )
    )
    styles.add(
        ParagraphStyle(
            name="CnCode",
            parent=styles["Code"],
            fontName=body_font,
            fontSize=8.8,
            leading=11.5,
            firstLineIndent=0,
            backColor=colors.HexColor("#F4F7FA"),
        )
    )
    styles.add(
        ParagraphStyle(
            name="CnCaption",
            parent=styles["Normal"],
            fontName=body_font,
            fontSize=8.8,
            leading=11,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#4D5D6C"),
            spaceAfter=5,
        )
    )

    story = []
    first_heading = True
    for el in elements:
        if el.kind == "heading":
            if first_heading:
                story.append(Paragraph(html_escape(el.text), styles["CnTitle"]))
                first_heading = False
            else:
                if el.level == 1:
                    story.append(PageBreak())
                story.append(Paragraph(html_escape(el.text), styles["CnH1" if el.level == 1 else "CnH2"]))
        elif el.kind == "paragraph":
            story.append(Paragraph(html_escape(el.text), styles["CnBody"]))
        elif el.kind == "bullets":
            items = [
                ListItem(Paragraph(html_escape(item), styles["CnBullet"]), bulletColor=colors.HexColor("#1F4E79"))
                for item in el.items
            ]
            story.append(ListFlowable(items, bulletType="bullet", leftIndent=12, bulletFontName=body_font))
            story.append(Spacer(1, 3))
        elif el.kind == "table":
            data = [
                [Paragraph(html_escape(cell), styles["CnTable"]) for cell in el.headers],
                *[[Paragraph(html_escape(cell), styles["CnTable"]) for cell in row] for row in el.rows],
            ]
            total_width = A4[0] - 4 * cm
            if el.widths:
                scale = total_width / sum(el.widths)
                col_widths = [w * scale for w in el.widths]
            else:
                col_widths = [total_width / len(el.headers)] * len(el.headers)
            table = Table(data, colWidths=col_widths, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, -1), body_font),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#A7B7C7")),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 5),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.append(KeepTogether([table, Spacer(1, 7)]))
        elif el.kind == "code":
            code_lines = [Paragraph(html_escape(line), styles["CnCode"]) for line in el.text.splitlines()]
            story.append(KeepTogether(code_lines + [Spacer(1, 5)]))
        elif el.kind == "image":
            image_path = resolve_asset(el.text)
            if image_path.exists():
                max_width = A4[0] - 4.0 * cm
                max_height = 9.6 * cm
                img = PdfImage(str(image_path))
                ratio = min(max_width / img.imageWidth, max_height / img.imageHeight)
                img.drawWidth = img.imageWidth * ratio
                img.drawHeight = img.imageHeight * ratio
                caption = Paragraph(html_escape(el.caption), styles["CnCaption"])
                story.append(KeepTogether([img, Spacer(1, 3), caption, Spacer(1, 7)]))
        elif el.kind == "pagebreak":
            story.append(PageBreak())

    def header_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(body_font, 8.5)
        canvas.setFillColor(colors.HexColor("#4D5D6C"))
        canvas.drawCentredString(A4[0] / 2, A4[1] - 1.15 * cm, f"{FULL_NAME} 软件设计说明书")
        canvas.drawCentredString(
            A4[0] / 2,
            0.95 * cm,
            f"著作权人：{COPYRIGHT_OWNER}    第 {doc.page} 页",
        )
        canvas.restoreState()

    path.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=2.0 * cm,
        rightMargin=2.0 * cm,
        topMargin=1.65 * cm,
        bottomMargin=1.55 * cm,
    )
    document.build(story, onFirstPage=header_footer, onLaterPages=header_footer)


def assert_formal_text(markdown: str) -> None:
    forbidden = (
        "排版说明",
        "关于著作权人署名",
        "请替换",
        "正式提交前优化计划",
        "SE-KB-Engine" + "知识库构建系统软件",
        "SE-KB-Engine " + "V0.1.0",
        "V0.1.0" + "",
        "0.1.0" + "",
        "内部发行版本号",
        "爬虫",
        "大规模抓取",
        "自动下载",
        "网页下载",
    )
    found = [word for word in forbidden if word in markdown]
    if found:
        raise ValueError(f"Formal document contains forbidden internal markers: {found}")
    if markdown.count(INTERNAL_CODE_NAME) != 1:
        raise ValueError(f"Expected exactly one internal code-name explanation for {INTERNAL_CODE_NAME}.")


def write_outputs() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    markdown = render_markdown(ELEMENTS)
    assert_formal_text(markdown)

    DOC_SOURCE.write_text(markdown, encoding="utf-8")
    (OUTPUT_DIR / f"{OUTPUT_BASENAME}.md").write_text(markdown, encoding="utf-8")
    INTERNAL_CHECKLIST.write_text(INTERNAL_CHECKLIST_TEXT, encoding="utf-8")

    render_docx(ELEMENTS, OUTPUT_DIR / f"{OUTPUT_BASENAME}.docx")
    render_pdf(ELEMENTS, OUTPUT_DIR / f"{OUTPUT_BASENAME}.pdf")


if __name__ == "__main__":
    write_outputs()
    print("Generated formal copyright design document:")
    print(DOC_SOURCE)
    print(OUTPUT_DIR / f"{OUTPUT_BASENAME}.docx")
    print(OUTPUT_DIR / f"{OUTPUT_BASENAME}.pdf")
